import io
import logging
from pathlib import Path
from typing import List, Dict, Any, Union
import pypdf

logger = logging.getLogger(__name__)

class DocumentParser:
    """Parses various document types (PDF, Word DOCX, Plain Text, Markdown) into structured items with metadata."""

    @staticmethod
    def parse_pdf(file_input: Union[str, Path, bytes, io.BytesIO], filename: str = "document.pdf") -> List[Dict[str, Any]]:
        """Parses a PDF file extracting text per page."""
        pages_data = []
        try:
            if isinstance(file_input, (str, Path)):
                reader = pypdf.PdfReader(str(file_input))
            else:
                if isinstance(file_input, bytes):
                    file_input = io.BytesIO(file_input)
                reader = pypdf.PdfReader(file_input)

            for page_idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                text = text.strip()
                if text:
                    pages_data.append({
                        "text": text,
                        "page_number": page_idx + 1,
                        "source": filename,
                        "doc_type": "pdf"
                    })
        except Exception as e:
            logger.error(f"Error parsing PDF {filename}: {e}")
        return pages_data

    @staticmethod
    def parse_docx(file_input: Union[str, Path, bytes, io.BytesIO], filename: str = "document.docx") -> List[Dict[str, Any]]:
        """Parses a Word (.docx) file extracting paragraphs and tables."""
        try:
            import docx
            if isinstance(file_input, bytes):
                file_input = io.BytesIO(file_input)
            
            doc = docx.Document(file_input)
            text_lines = []

            # Extract paragraphs
            for p in doc.paragraphs:
                clean_p = p.text.strip()
                if clean_p:
                    text_lines.append(clean_p)

            # Extract tables (common in lab and patient reports)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_cells:
                        text_lines.append(" | ".join(row_cells))

            full_text = "\n".join(text_lines)
            if full_text.strip():
                return [{
                    "text": full_text.strip(),
                    "page_number": 1,
                    "source": filename,
                    "doc_type": "docx"
                }]
        except Exception as e:
            logger.error(f"Error parsing Word DOCX {filename}: {e}")
        return []

    @staticmethod
    def parse_text(text_content: str, filename: str = "document.txt") -> List[Dict[str, Any]]:
        """Parses raw text or markdown content."""
        clean_text = text_content.strip()
        if not clean_text:
            return []
        return [{
            "text": clean_text,
            "page_number": 1,
            "source": filename,
            "doc_type": "text"
        }]

    @classmethod
    def parse_file(cls, file_path_or_bytes: Union[str, Path, bytes], filename: str = "") -> List[Dict[str, Any]]:
        """Auto-detects file type (PDF, Word DOCX, TXT, MD) and parses accordingly."""
        if isinstance(file_path_or_bytes, (str, Path)):
            path = Path(file_path_or_bytes)
            name = filename or path.name
            suffix = path.suffix.lower()
            if suffix == ".pdf":
                return cls.parse_pdf(path, filename=name)
            elif suffix in [".docx", ".doc"]:
                return cls.parse_docx(path, filename=name)
            else:
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    return cls.parse_text(f.read(), filename=name)
        else:
            # Bytes input (e.g. from web upload)
            name = filename or "uploaded_doc"
            suffix = Path(name).suffix.lower()
            if suffix == ".pdf":
                return cls.parse_pdf(file_path_or_bytes, filename=name)
            elif suffix in [".docx", ".doc"]:
                return cls.parse_docx(file_path_or_bytes, filename=name)
            else:
                try:
                    text = file_path_or_bytes.decode("utf-8", errors="ignore")
                except Exception:
                    text = str(file_path_or_bytes)
                return cls.parse_text(text, filename=name)
